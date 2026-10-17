from datetime import date
import io
from django.db import IntegrityError, models, transaction
from django.db.models import F, Q, Prefetch, ProtectedError, RestrictedError
from django.http import HttpResponse, JsonResponse
from django.shortcuts import render, redirect, get_object_or_404
from django.contrib.auth.decorators import login_required, user_passes_test
from django.urls import reverse
from django.views.decorators.http import require_POST
from django.contrib import messages
from django.utils import timezone

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from reportlab.lib import colors
from reportlab.lib.pagesizes import letter, landscape, portrait
from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle

from core.models import AcademicSession
from academics.models import AcademicClass, Section, Stream, Subject
from .models import Student, StudentEnrollment, StudentSubjectEnrollment
from .forms import StudentEnrollmentForm
from .utils import mask_contact, is_teacher_or_staff


from datetime import date
from django.shortcuts import render, redirect, get_object_or_404
from django.urls import reverse
from django.contrib import messages
from django.db import transaction, IntegrityError
from django.db.models import Prefetch, Sum, Q
from django.contrib.auth.decorators import login_required, user_passes_test

# Adjust this import to match where your models are defined
# from .models import AcademicSession, AcademicClass, Section, Stream, StudentEnrollment


def is_staff_member(user):
    return user.is_authenticated and (user.is_staff or hasattr(user, 'staff_profile'))


# =====================================================================
# STUDENT ROSTER & DIRECTORY
# =====================================================================

@login_required
def student_list_view(request):
    sessions = AcademicSession.objects.all().order_by('-start_date')
    selected_session_id = request.GET.get('session', '').strip()

    # 1. Resolve selected session safely by ID or fallback
    selected_session = None
    if selected_session_id:
        try:
            selected_session = sessions.filter(id=int(selected_session_id)).first()
        except (ValueError, TypeError):
            selected_session = None

    if not selected_session:
        selected_session = sessions.filter(is_current=True).first() or sessions.first()

    # Query parameters
    selected_grade = request.GET.get('grade', '').strip()
    selected_section = request.GET.get('section', '').strip()
    selected_status = request.GET.get('status', '').strip()
    query = request.GET.get('q', '').strip()

    # 2. Base Queryset (Use session_id to avoid cross-app model instance type collisions)
    if selected_session:
        enrollments = StudentEnrollment.objects.filter(
            session_id=selected_session.id
        ).select_related(
            'student',
            'section',
            'section__academic_class',
            'section__stream',
            'session'
        )
        sec_qs = Section.objects.filter(
            session_id=selected_session.id
        ).select_related(
            'academic_class', 
            'stream'
        ).order_by('academic_class__numeric_level', 'name')
    else:
        enrollments = StudentEnrollment.objects.none()
        sec_qs = Section.objects.none()

    # 3. Filter: Grade / Class
    if selected_grade:
        enrollments = enrollments.filter(section__academic_class__numeric_level=selected_grade)
        sec_qs = sec_qs.filter(academic_class__numeric_level=selected_grade)

    # 4. Filter: Specific Section
    if selected_section:
        try:
            enrollments = enrollments.filter(section_id=int(selected_section))
        except (ValueError, TypeError):
            pass

    # 5. Filter: Informed Status
    if selected_status == 'informed':
        enrollments = enrollments.filter(is_informed=True)
    elif selected_status == 'pending':
        enrollments = enrollments.filter(is_informed=False)

    # 6. Filter: Search Keyword
    if query:
        enrollments = enrollments.filter(
            Q(student__student_name__icontains=query) |
            Q(student__admission_number__icontains=query) |
            Q(roll_number__icontains=query) |
            Q(student__father_name__icontains=query) |
            Q(student__mobile_number__icontains=query)
        )

    # 7. Access permission flag
    teacher_access = request.user.is_staff or request.user.is_superuser

    # 8. Build student record items
    records = []
    for enr in enrollments:
        mobile = getattr(enr.student, 'mobile_number', '') or ''
        records.append({
            'student': enr.student,
            'enrollment': enr,
            'phone': mobile,
            'raw_phone': mobile,
            'can_access': teacher_access,
        })

    # Summary metrics
    grades = AcademicClass.objects.all().order_by('numeric_level')
    total_students = enrollments.count()
    total_marked = enrollments.filter(is_informed=True).count()
    total_remaining = max(0, total_students - total_marked)

    return render(request, 'students/student_list.html', {
        'sessions': sessions,
        'selected_session': selected_session,
        'grades': grades,
        'available_sections': sec_qs,
        'selected_grade': selected_grade,
        'selected_section': selected_section,
        'selected_status': selected_status,
        'query': query,
        'records': records,
        'total_students': total_students,
        'total_marked': total_marked,
        'total_remaining': total_remaining,
        'teacher_access': teacher_access,
        'has_active_session': selected_session is not None,
    })


@login_required
@require_POST
def toggle_student_informed(request, enrollment_id):
    """Toggle or set informed status for a single student enrollment."""
    enrollment = get_object_or_404(StudentEnrollment, id=enrollment_id)
    action = request.POST.get('action')

    if action == 'mark':
        enrollment.is_informed = True
        enrollment.informed_at = timezone.now()
    elif action == 'clear':
        enrollment.is_informed = False
        enrollment.informed_at = None
    else:
        enrollment.is_informed = not enrollment.is_informed
        enrollment.informed_at = timezone.now() if enrollment.is_informed else None

    enrollment.save(update_fields=['is_informed', 'informed_at'])
    return JsonResponse({
        'status': 'success',
        'is_informed': enrollment.is_informed,
        'enrollment_id': enrollment.id
    })

# Inside manage_academic_structure POST block or in a dedicated view:
def toggle_session_flags(request):
    if request.method == 'POST':
        action = request.POST.get('action')
        session_id = request.POST.get('session_id')
        sess = get_object_or_404(AcademicSession, id=session_id)

        # 1. SET AS ACTIVE SESSION
        if action == 'set_active_session':
            with transaction.atomic():
                AcademicSession.objects.exclude(id=sess.id).update(is_current=False)
                sess.is_current = True
                sess.save(update_fields=['is_current'])
            messages.success(request, f"Session '{sess.name}' is now marked as the ACTIVE academic session.")

        # 2. TOGGLE MODIFICATION LOCK (ENABLE / DISABLE)
        elif action == 'toggle_modification_lock':
            sess.allow_modifications = not sess.allow_modifications
            sess.save(update_fields=['allow_modifications'])
            status_text = "ENABLED" if sess.allow_modifications else "LOCKED (Read-only)"
            messages.success(request, f"Student modifications are now {status_text} for session '{sess.name}'.")

        # Redirect back to where the user came from
        next_url = request.POST.get('next') or request.META.get('HTTP_REFERER', '/')
        return redirect(next_url)

    

@login_required
@require_POST
def clear_section_informed(request, section_id):
    """Clear all 'informed' marks for students in a specific section and session."""
    session_id = request.POST.get('session_id')
    qs = StudentEnrollment.objects.filter(section_id=section_id)

    if session_id:
        qs = qs.filter(session_id=session_id)

    updated_count = qs.filter(is_informed=True).update(is_informed=False, informed_at=None)

    return JsonResponse({
        'status': 'success',
        'cleared_count': updated_count
    })


def student_detail(request, pk):
    enrollments_prefetch = Prefetch(
        'enrollments',
        queryset=StudentEnrollment.objects.select_related(
            'session', 
            'section__academic_class', 
            'section__stream'
        ).prefetch_related(
            Prefetch(
                'enrolled_subjects',
                queryset=StudentSubjectEnrollment.objects.select_related('subject')
            )
        ).order_by('-session__start_date', 'roll_number')
    )

    student = get_object_or_404(
        Student.objects.prefetch_related(enrollments_prefetch),
        pk=pk
    )

    return render(request, 'students/student_detail.html', {
        'student': student,
    })


@login_required
def student_create(request):
    active_session = AcademicSession.objects.filter(is_current=True).first()

    if not active_session:
        messages.error(request, "No active academic session found. Please activate one in Academic Settings.")
        return redirect('student_list')

    if request.method == 'POST':
        form = StudentEnrollmentForm(request.POST, request.FILES)
        if form.is_valid():
            student = form.save()
            messages.success(request, f"Student {student.student_name} enrolled successfully!")
            return redirect('student_list')
    else:
        form = StudentEnrollmentForm()

    return render(request, 'students/student_form.html', {
        'form': form,
        'active_session': active_session,
        'is_edit': False,
    })


@login_required
@user_passes_test(is_teacher_or_staff, login_url='login')
def student_update(request, pk):
    student = get_object_or_404(Student, pk=pk)
    active_session = AcademicSession.objects.filter(is_current=True).first()

    enrollment = (
        student.enrollments.filter(session=active_session).select_related('section', 'session').first()
        or student.enrollments.select_related('section', 'session').first()
    )

    if request.method == 'POST':
        form = StudentEnrollmentForm(request.POST, request.FILES, instance=student)
        if form.is_valid():
            form.save()
            messages.success(request, f"Updated profile for {student.student_name}.")
            return redirect('student_detail', pk=student.pk)
    else:
        form = StudentEnrollmentForm(instance=student)

    enrolled_subject_ids = []
    if enrollment:
        enrolled_subject_ids = list(
            enrollment.enrolled_subjects.values_list('subject_id', flat=True)
        )

    sessions = AcademicSession.objects.all().order_by('-start_date')
    sections = Section.objects.select_related('academic_class', 'stream').all()

    latest_enrollment = student.enrollments.order_by('-session__start_date').first()
    if latest_enrollment and not latest_enrollment.session.allow_modifications:
        messages.error(request, f"Modifications are locked for Academic Session '{latest_enrollment.session.name}'.")
        return redirect('student_list')
    
    return render(request, 'students/student_form.html', {
        'form': form,
        'student': student,
        'enrollment': enrollment,
        'sessions': sessions,
        'sections': sections,
        'active_session': active_session,
        'enrolled_subject_ids': enrolled_subject_ids,
        'is_edit': True,
    })


@login_required
@user_passes_test(is_teacher_or_staff, login_url='login')
@require_POST
def student_delete(request, pk):
    student = get_object_or_404(Student, pk=pk)
    name = student.student_name
    student.delete()
    messages.success(request, f"Student {name} deleted successfully.")
    return redirect('student_list')


# SINGLE, CLEAN SUBJECT-BY-SECTION API (No duplicates)
def get_subjects_by_section(request):
    section_id = request.GET.get('section_id')
    if not section_id:
        return JsonResponse({'subjects': [], 'error': 'No section_id provided'})

    try:
        section = Section.objects.select_related('academic_class', 'stream').get(id=section_id)
    except Section.DoesNotExist:
        return JsonResponse({'subjects': [], 'error': f'Section {section_id} not found'})

    qs = Subject.objects.filter(academic_class=section.academic_class)

    has_stream_attr = any(f.name == 'stream' for f in Subject._meta.get_fields())
    if has_stream_attr:
        if section.stream:
            qs = qs.filter(Q(stream=section.stream) | Q(stream__isnull=True))
        else:
            qs = qs.filter(stream__isnull=True)

    data = []
    for s in qs.order_by('name'):
        stype = 'General'
        if hasattr(s, 'get_subject_type_display'):
            stype = s.get_subject_type_display()
        elif hasattr(s, 'subject_type') and s.subject_type:
            stype = str(s.subject_type).title()

        data.append({
            'id': s.id,
            'name': getattr(s, 'name', f'Subject #{s.id}'),
            'code': getattr(s, 'code', '') or '',
            'type': stype,
        })

    return JsonResponse({'subjects': data, 'count': len(data)})


# =====================================================================
# SESSION PROMOTION & ROLLOVER
# =====================================================================

@login_required
@user_passes_test(is_teacher_or_staff, login_url='login')
def promote_session(request):
    sessions = AcademicSession.objects.all().order_by('-start_date')
    active_session = AcademicSession.objects.filter(is_current=True).first() or sessions.first()
    sections = Section.objects.select_related('academic_class', 'stream').all()

    from_section_id = request.GET.get('from_section')
    from_section = None
    active_students = StudentEnrollment.objects.none()

    if from_section_id:
        from_section = sections.filter(id=from_section_id).first()
        if from_section and active_session:
            active_students = StudentEnrollment.objects.filter(
                session=active_session,
                section=from_section,
                status='active'
            ).select_related('student', 'section__academic_class', 'section__stream').order_by('roll_number')

    if request.method == 'POST':
        from_section_id = request.POST.get('from_section')
        to_section_id = request.POST.get('to_section')
        target_session_id = request.POST.get('target_session')
        selected_student_ids = request.POST.getlist('student_ids')

        if not (from_section_id and to_section_id and target_session_id):
            messages.error(request, "Please select source section, target section, and target session.")
            return redirect('promote_session')

        if not selected_student_ids:
            messages.error(request, "Please select at least one student to promote.")
            return redirect(f"{request.path}?from_section={from_section_id}")

        with transaction.atomic():
            for s_id in selected_student_ids:
                old_enr = StudentEnrollment.objects.filter(
                    student_id=s_id, 
                    section_id=from_section_id,
                    session=active_session
                ).first()
                if old_enr:
                    old_enr.status = 'promoted'
                    old_enr.save()

                StudentEnrollment.objects.get_or_create(
                    student_id=s_id,
                    session_id=target_session_id,
                    section_id=to_section_id,
                    defaults={
                        'roll_number': old_enr.roll_number if old_enr else 1,
                        'status': 'active'
                    }
                )

        messages.success(request, f"Successfully promoted {len(selected_student_ids)} students.")
        return redirect('student_list')

    return render(request, 'students/promote_session.html', {
        'sessions': sessions,
        'active_session': active_session,
        'sections': sections,
        'from_section': from_section,
        'from_section_id': from_section_id,
        'active_students': active_students,
    })


# =====================================================================
# ACADEMIC STRUCTURE CONTROLLER (FULL MODAL & SEED ACTIONS HANDLED)
# =====================================================================


# @login_required
# @user_passes_test(is_staff_member, login_url='login')
# def manage_academic_structure_view(request):
#     # 1. Resolve Active Session (create default if database has none)
#     active_session = AcademicSession.objects.filter(is_current=True).first()
#     if not active_session:
#         active_session = AcademicSession.objects.first()
#         if not active_session:
#             active_session = AcademicSession.objects.create(name="2026-2027", is_current=True)
#         else:
#             active_session.is_current = True
#             active_session.save(update_fields=['is_current'])

#     if request.method == 'POST':
#         action = request.POST.get('action')

#         # =====================================================================
#         # 1. SEED DEFAULT STRUCTURE (9th-3Sec, 10th-3Sec, 11th/12th Sci & Arts)
#         # =====================================================================
#         if action == 'seed_default_structure':
#             with transaction.atomic():
#                 # A. Permanent Classes
#                 c9, _ = AcademicClass.objects.get_or_create(
#                     numeric_level=9,
#                     defaults={'name': 'Class IX (Secondary Foundation)', 'is_intermediate': False}
#                 )
#                 c10, _ = AcademicClass.objects.get_or_create(
#                     numeric_level=10,
#                     defaults={'name': 'Class X (Secondary Board)', 'is_intermediate': False}
#                 )
#                 c11, _ = AcademicClass.objects.get_or_create(
#                     numeric_level=11,
#                     defaults={'name': 'Class XI (Senior Secondary)', 'is_intermediate': True}
#                 )
#                 c12, _ = AcademicClass.objects.get_or_create(
#                     numeric_level=12,
#                     defaults={'name': 'Class XII (Senior Secondary Terminal)', 'is_intermediate': True}
#                 )

#                 # B. Faculty Streams
#                 stream_sci, _ = Stream.objects.get_or_create(
#                     code='SCIENCE',
#                     defaults={
#                         'name': 'Faculty of Science',
#                         'faculty': 'SCIENCE',
#                         'sanctioned_seats': 60,
#                         'description': 'Physics, Chemistry, Mathematics, Biology, English, Hindi'
#                     }
#                 )
#                 stream_arts, _ = Stream.objects.get_or_create(
#                     code='ARTS',
#                     defaults={
#                         'name': 'Faculty of Arts',
#                         'faculty': 'ARTS',
#                         'sanctioned_seats': 60,
#                         'description': 'History, Political Science, Geography, Economics, Hindi, English'
#                     }
#                 )

#                 # Remove legacy/placeholder Sec A and Sec B from 11th & 12th if present
#                 Section.objects.filter(academic_class__in=[c11, c12], name__in=['A', 'B']).delete()

#                 # C. Deploy sections across all sessions
#                 target_sessions = list(AcademicSession.objects.all())
#                 if not target_sessions:
#                     target_sessions = [active_session]

#                 for sess in target_sessions:
#                     # Class 9 -> 3 Sections (A, B, C)
#                     for sec in ['A', 'B', 'C']:
#                         Section.objects.get_or_create(
#                             academic_class=c9,
#                             session=sess,
#                             name=sec,
#                             defaults={'capacity': 40, 'room_number': f'9-{sec}'}
#                         )

#                     # Class 10 -> 3 Sections (A, B, C)
#                     for sec in ['A', 'B', 'C']:
#                         Section.objects.get_or_create(
#                             academic_class=c10,
#                             session=sess,
#                             name=sec,
#                             defaults={'capacity': 40, 'room_number': f'10-{sec}'}
#                         )

#                     # Class 11 -> Science & Arts
#                     Section.objects.get_or_create(
#                         academic_class=c11,
#                         session=sess,
#                         name='Science',
#                         defaults={'stream': stream_sci, 'capacity': 60, 'room_number': 'Lab 1'}
#                     )
#                     Section.objects.get_or_create(
#                         academic_class=c11,
#                         session=sess,
#                         name='Arts',
#                         defaults={'stream': stream_arts, 'capacity': 60, 'room_number': 'Hall 1'}
#                     )

#                     # Class 12 -> Science & Arts
#                     Section.objects.get_or_create(
#                         academic_class=c12,
#                         session=sess,
#                         name='Science',
#                         defaults={'stream': stream_sci, 'capacity': 60, 'room_number': 'Lab 2'}
#                     )
#                     Section.objects.get_or_create(
#                         academic_class=c12,
#                         session=sess,
#                         name='Arts',
#                         defaults={'stream': stream_arts, 'capacity': 60, 'room_number': 'Hall 2'}
#                     )

#             messages.success(request, "Academic structure successfully seeded: Class 9 (A, B, C), Class 10 (A, B, C), and Classes 11 & 12 (Science & Arts).")
#             return redirect('manage_academic_structure')

#         # =====================================================================
#         # 2. SESSION ACTIONS
#         # =====================================================================
#         elif action == 'create_session':
#             s_name = request.POST.get('session_name', '').strip()
#             s_date = request.POST.get('start_date')
#             e_date = request.POST.get('end_date')
#             if s_name and s_date and e_date:
#                 AcademicSession.objects.create(name=s_name, start_date=s_date, end_date=e_date, is_current=False)
#                 messages.success(request, f"Academic Session '{s_name}' created.")
#             return redirect('manage_academic_structure')

#         # =====================================================================
#         # 3. CLASS ACTIONS
#         # =====================================================================
#         elif action == 'create_class':
#             lvl = request.POST.get('numeric_level')
#             name = request.POST.get('name', '').strip()
#             is_inter = request.POST.get('is_intermediate') == 'on'
#             if lvl and name:
#                 AcademicClass.objects.get_or_create(
#                     numeric_level=int(lvl),
#                     defaults={'name': name, 'is_intermediate': is_inter}
#                 )
#                 messages.success(request, f"Class '{name}' created.")
#             return redirect('manage_academic_structure')

#         elif action == 'update_class':
#             cid = request.POST.get('class_id')
#             c = get_object_or_404(AcademicClass, id=cid)
#             c.numeric_level = int(request.POST.get('numeric_level', c.numeric_level))
#             c.name = request.POST.get('name', c.name).strip()
#             c.is_intermediate = request.POST.get('is_intermediate') == 'on'
#             c.save()
#             messages.success(request, f"Class '{c.name}' updated.")
#             return redirect('manage_academic_structure')

#         elif action == 'delete_class':
#             cid = request.POST.get('target_id')
#             c = get_object_or_404(AcademicClass, id=cid)
#             c.delete()
#             messages.success(request, "Class removed.")
#             return redirect('manage_academic_structure')

#         # =====================================================================
#         # 4. SECTION ACTIONS
#         # =====================================================================
#         elif action == 'create_section':
#             cid = request.POST.get('class_id')
#             sec_name = request.POST.get('section_name', '').strip()
#             stream_id = request.POST.get('stream_id')
#             cap = int(request.POST.get('capacity', 40) or 40)
#             room = request.POST.get('room_number', '').strip()

#             ac_class = get_object_or_404(AcademicClass, id=cid)
#             stream = Stream.objects.filter(id=stream_id).first() if stream_id else None

#             if Section.objects.filter(academic_class=ac_class, name__iexact=sec_name, session=active_session).exists():
#                 messages.error(request, f"Section '{sec_name}' already exists in {ac_class.name}.")
#             else:
#                 Section.objects.create(
#                     session=active_session,
#                     academic_class=ac_class,
#                     name=sec_name,
#                     stream=stream,
#                     capacity=cap,
#                     room_number=room
#                 )
#                 messages.success(request, f"Section '{sec_name}' registered under {ac_class.name}.")
#             return redirect('manage_academic_structure')

#         elif action == 'update_section':
#             sec_id = request.POST.get('section_id')
#             sec = get_object_or_404(Section, id=sec_id)
#             sec.academic_class_id = request.POST.get('class_id', sec.academic_class_id)
#             sec.name = request.POST.get('section_name', sec.name).strip()
#             stream_id = request.POST.get('stream_id')
#             sec.stream = Stream.objects.filter(id=stream_id).first() if stream_id else None
#             sec.capacity = int(request.POST.get('capacity', sec.capacity) or 40)
#             sec.room_number = request.POST.get('room_number', sec.room_number).strip()
#             sec.save()
#             messages.success(request, f"Section '{sec.name}' updated.")
#             return redirect('manage_academic_structure')

#         elif action == 'delete_section':
#             sec_id = request.POST.get('target_id')
#             sec = get_object_or_404(Section, id=sec_id)
#             sec.delete()
#             messages.success(request, "Section removed.")
#             return redirect('manage_academic_structure')

#         # =====================================================================
#         # 5. STREAM ACTIONS
#         # =====================================================================
#         elif action == 'create_stream':
#             sname = request.POST.get('stream_name', '').strip()
#             code = request.POST.get('stream_code', '').strip().upper()
#             faculty = request.POST.get('faculty', 'SCIENCE')
#             seats = int(request.POST.get('sanctioned_seats', 60) or 60)
#             desc = request.POST.get('stream_desc', '').strip()
#             if sname and code:
#                 Stream.objects.create(
#                     name=sname, code=code, faculty=faculty,
#                     sanctioned_seats=seats, description=desc
#                 )
#                 messages.success(request, f"Stream '{sname}' registered.")
#             return redirect('manage_academic_structure')

#         elif action == 'update_stream':
#             sid = request.POST.get('stream_id')
#             st = get_object_or_404(Stream, id=sid)
#             st.name = request.POST.get('stream_name', st.name).strip()
#             st.code = request.POST.get('stream_code', st.code).strip().upper()
#             st.faculty = request.POST.get('faculty', st.faculty)
#             st.sanctioned_seats = int(request.POST.get('sanctioned_seats', st.sanctioned_seats) or 60)
#             st.description = request.POST.get('stream_desc', st.description).strip()
#             st.save()
#             messages.success(request, f"Stream '{st.name}' updated.")
#             return redirect('manage_academic_structure')

#         elif action == 'delete_stream':
#             sid = request.POST.get('target_id')
#             st = get_object_or_404(Stream, id=sid)
#             st.delete()
#             messages.success(request, "Stream removed.")
#             return redirect('manage_academic_structure')

#         # =====================================================================
#         # 6. TRANSFER / PROMOTION MODAL ACTION
#         # =====================================================================
#         elif action == 'transfer_student_grade':
#             stu_id = request.POST.get('student_id')
#             cur_enr_id = request.POST.get('current_enrollment_id')
#             target_sess_id = request.POST.get('target_session_id')
#             target_sec_id = request.POST.get('target_section_id')
#             new_roll = int(request.POST.get('new_roll_number', 1) or 1)

#             with transaction.atomic():
#                 cur_enr = StudentEnrollment.objects.filter(id=cur_enr_id).first()
#                 if cur_enr:
#                     cur_enr.status = 'promoted'
#                     cur_enr.save(update_fields=['status'])

#                 StudentEnrollment.objects.update_or_create(
#                     student_id=stu_id,
#                     session_id=target_sess_id,
#                     defaults={'section_id': target_sec_id, 'roll_number': new_roll, 'status': 'active'}
#                 )
#             messages.success(request, "Student enrollment transferred/promoted.")
#             return redirect('manage_academic_structure')

#     # =========================================================================
#     # GET: PREPARE AND RENDER VIEW
#     # =========================================================================
#     sessions = AcademicSession.objects.all().order_by('-start_date')
#     selected_sess_id = request.GET.get('session')
#     if selected_sess_id:
#         view_session = sessions.filter(id=selected_sess_id).first() or active_session
#     else:
#         view_session = active_session

#     classes = AcademicClass.objects.prefetch_related(
#         Prefetch(
#             'sections',
#             queryset=Section.objects.filter(session=view_session).select_related('stream').order_by('name')
#         )
#     ).all().order_by('numeric_level')

#     streams = Stream.objects.all().order_by('faculty', 'code')

#     science_seats = sum(s.sanctioned_seats for s in streams.filter(faculty='SCIENCE')) or 60
#     arts_seats = sum(s.sanctioned_seats for s in streams.filter(faculty='ARTS')) or 60
#     total_intermediate_capacity = sum(
#         sec.capacity for c in classes if c.is_intermediate for sec in c.sections.all()
#     )

#     return render(request, 'students/manage_academic_structure.html', {
#         'classes': classes,
#         'streams': streams,
#         'sessions': sessions,
#         'active_session': view_session,
#         'science_seats': science_seats,
#         'arts_seats': arts_seats,
#         'total_intermediate_capacity': total_intermediate_capacity,
#     })

def can_manage_academic_structure(user):
    return user.is_authenticated and (user.is_staff or user.is_superuser)


def seed_session_structure(session_obj):
    """Creates standard Classes (9-12), Streams, and Sections for session_obj."""
    if not session_obj:
        return

    # 1. Create Streams
    stream_sci, _ = Stream.objects.get_or_create(
        code='SCIENCE',
        defaults={
            'name': 'Science (PCM/PCB)',
            'faculty': 'SCIENCE',
            'sanctioned_seats': 60,
            'description': 'Physics, Chemistry, Mathematics, Biology',
        }
    )
    stream_arts, _ = Stream.objects.get_or_create(
        code='ARTS',
        defaults={
            'name': 'Arts (I.A)',
            'faculty': 'ARTS',
            'sanctioned_seats': 60,
            'description': 'History, Political Science, Geography, Economics',
        }
    )

    # 2. Create Classes
    classes_config = [
        (9, 'Class IX (Secondary Foundation)', False),
        (10, 'Class X (Secondary Board)', False),
        (11, 'Class XI (Senior Secondary)', True),
        (12, 'Class XII (Senior Secondary Terminal)', True),
    ]

    for num, name, is_inter in classes_config:
        cls_obj, _ = AcademicClass.objects.get_or_create(
            numeric_level=num,
            defaults={'name': name, 'is_intermediate': is_inter}
        )

        if is_inter:
            # For 11th & 12th, give them unique names ('Science' and 'Arts')
            # to guarantee no database unique_together crashes!
            Section.objects.get_or_create(
                session=session_obj,
                academic_class=cls_obj,
                name='Science',
                defaults={
                    'stream': stream_sci,
                    'capacity': 60,
                }
            )
            Section.objects.get_or_create(
                session=session_obj,
                academic_class=cls_obj,
                name='Arts',
                defaults={
                    'stream': stream_arts,
                    'capacity': 60,
                }
            )
        else:
            # For 9th & 10th
            Section.objects.get_or_create(
                session=session_obj,
                academic_class=cls_obj,
                name='A',
                defaults={'capacity': 40, 'stream': None}
            )
            Section.objects.get_or_create(
                session=session_obj,
                academic_class=cls_obj,
                name='B',
                defaults={'capacity': 40, 'stream': None}
            )


@login_required
@user_passes_test(can_manage_academic_structure, login_url='/admin/login/')
def manage_academic_structure(request):
    session_param = request.GET.get('session') or request.POST.get('session_id') or ''

    def redirect_with_session(target_id=None):
        base_url = reverse('manage_academic_structure')
        sid = target_id or session_param
        if sid:
            return redirect(f"{base_url}?session={sid}")
        return redirect('manage_academic_structure')

    # Handle POST
    if request.method == 'POST':
        action = request.POST.get('action')

        if action == 'seed_default_structure':
            target_sess_id = request.POST.get('session_id') or session_param
            target_sess = AcademicSession.objects.filter(id=target_sess_id).first() if target_sess_id else None

            if not target_sess:
                target_sess = AcademicSession.objects.filter(is_current=True).first() or AcademicSession.objects.first()

            if not target_sess:
                target_sess = AcademicSession.objects.create(
                    name='2026-2027',
                    start_date=date(2026, 4, 1),
                    end_date=date(2027, 3, 31),
                    is_current=True
                )

            with transaction.atomic():
                seed_session_structure(target_sess)

            messages.success(request, f"Standard structure created for {target_sess.name}!")
            return redirect_with_session(target_sess.id)

        elif action == 'create_session':
            name = request.POST.get('session_name', '').strip()
            start_date = request.POST.get('start_date')
            end_date = request.POST.get('end_date')

            if name and start_date and end_date:
                try:
                    with transaction.atomic():
                        new_session = AcademicSession.objects.create(
                            name=name,
                            start_date=start_date,
                            end_date=end_date
                        )
                        seed_session_structure(new_session)
                    messages.success(request, f"Session '{name}' created with default structure!")
                    return redirect_with_session(new_session.id)
                except IntegrityError:
                    messages.error(request, f"Session '{name}' already exists.")
            return redirect_with_session()

        elif action == 'create_section':
            class_id = request.POST.get('class_id')
            session_id = request.POST.get('session_id') or session_param
            sec_name = request.POST.get('section_name', '').strip().upper()
            stream_id = request.POST.get('stream_id') or None
            room_number = request.POST.get('room_number', '').strip()
            capacity = int(request.POST.get('capacity') or 40)

            try:
                Section.objects.create(
                    session_id=session_id,
                    academic_class_id=class_id,
                    name=sec_name,
                    stream_id=stream_id,
                    room_number=room_number,
                    capacity=capacity
                )
                messages.success(request, f"Section '{sec_name}' created.")
            except IntegrityError:
                messages.error(request, "Section already exists.")
            return redirect_with_session()

        elif action == 'update_section':
            sec = get_object_or_404(Section, id=request.POST.get('section_id'))
            sec.academic_class_id = request.POST.get('class_id') or sec.academic_class_id
            sec.name = request.POST.get('section_name', '').strip().upper()
            sec.stream_id = request.POST.get('stream_id') or None
            sec.room_number = request.POST.get('room_number', '').strip()
            sec.capacity = int(request.POST.get('capacity') or sec.capacity)
            sec.save()
            messages.success(request, f"Section '{sec.name}' updated.")
            return redirect_with_session()

        elif action == 'delete_section':
            sec = get_object_or_404(Section, id=request.POST.get('target_id'))
            sec.delete()
            messages.success(request, "Section deleted.")
            return redirect_with_session()

        elif action == 'create_stream':
            Stream.objects.create(
                name=request.POST.get('stream_name', '').strip(),
                code=request.POST.get('stream_code', '').strip().upper(),
                faculty=request.POST.get('faculty', 'SCIENCE'),
                sanctioned_seats=int(request.POST.get('sanctioned_seats') or 60),
                description=request.POST.get('stream_desc', '').strip()
            )
            messages.success(request, "Stream created.")
            return redirect_with_session()

        elif action == 'update_stream':
            st = get_object_or_404(Stream, id=request.POST.get('stream_id'))
            st.name = request.POST.get('stream_name', '').strip()
            st.code = request.POST.get('stream_code', '').strip().upper()
            st.faculty = request.POST.get('faculty', 'SCIENCE')
            st.sanctioned_seats = int(request.POST.get('sanctioned_seats') or st.sanctioned_seats)
            st.description = request.POST.get('stream_desc', '').strip()
            st.save()
            messages.success(request, f"Stream '{st.name}' updated.")
            return redirect_with_session()

        elif action == 'delete_stream':
            st = get_object_or_404(Stream, id=request.POST.get('target_id'))
            st.delete()
            messages.success(request, "Stream removed.")
            return redirect_with_session()

    # GET REQUEST
    sessions = AcademicSession.objects.all().order_by('-start_date')
    selected_session_id = request.GET.get('session')

    active_session = None
    if selected_session_id:
        active_session = sessions.filter(id=selected_session_id).first()
    if not active_session:
        active_session = sessions.filter(is_current=True).first() or sessions.first()

    # Force auto-provisioning if active session currently has 0 sections
    if active_session and not Section.objects.filter(session=active_session).exists():
        with transaction.atomic():
            seed_session_structure(active_session)

    # Attach sections directly to each class for the active session
    all_classes = list(AcademicClass.objects.all().order_by('numeric_level'))
    active_sections = list(
        Section.objects.filter(session=active_session).select_related('stream')
    ) if active_session else []

    # Bind sections cleanly so both .section_list and count work without template relation issues
    for c in all_classes:
        c.section_list = [s for s in active_sections if s.academic_class_id == c.id]
        c.section_count = len(c.section_list)

    streams = Stream.objects.all()
    science_seats = streams.filter(faculty='SCIENCE').aggregate(s=Sum('sanctioned_seats'))['s'] or 0
    arts_seats = streams.filter(faculty='ARTS').aggregate(s=Sum('sanctioned_seats'))['s'] or 0

    inter_sec_filter = Q(academic_class__is_intermediate=True)
    if active_session:
        inter_sec_filter &= Q(session=active_session)
    total_intermediate_capacity = Section.objects.filter(inter_sec_filter).aggregate(c=Sum('capacity'))['c'] or 0

    active_sections = list(
        Section.objects.filter(session=active_session).select_related('stream')
    ) if active_session else []

    total_sections_count = len(active_sections)

    return render(request, 'academics/manage_academic_structure.html', {
        'classes': all_classes,
        'streams': streams,
        'sessions': sessions,
        'active_session': active_session,
        'total_sections_count': total_sections_count,
        'science_seats': science_seats,
        'arts_seats': arts_seats,
        'total_intermediate_capacity': total_intermediate_capacity,
    })




@login_required
@user_passes_test(is_staff_member, login_url='login')
def edit_section_view(request, pk):
    section = get_object_or_404(Section, pk=pk)
    if request.method == 'POST':
        name = request.POST.get('section_name', '').strip().upper()
        stream_id = request.POST.get('stream_id')
        capacity = request.POST.get('capacity', 40)
        room = request.POST.get('room_number', '').strip()

        conflict = Section.objects.filter(
            academic_class=section.academic_class, 
            name=name, 
            session=section.session
        ).exclude(pk=section.pk).exists()

        if conflict:
            messages.error(request, f"Section {name} already exists in {section.academic_class.name}.")
        else:
            section.name = name
            section.stream = Stream.objects.filter(id=stream_id).first() if stream_id else None
            section.capacity = int(capacity or 40)
            section.room_number = room
            section.save()
            messages.success(request, f"Section {section.name} updated.")

    return redirect('manage_academic_structure')


@login_required
@user_passes_test(is_staff_member, login_url='login')
def delete_section_view(request, pk):
    section = get_object_or_404(Section, pk=pk)
    try:
        if section.enrollments.exists():
            messages.error(request, f"Cannot delete Section {section.name} because students are currently enrolled.")
        else:
            section.delete()
            messages.success(request, f"Section {section.name} deleted.")
    except (ProtectedError, RestrictedError):
        messages.error(request, f"Cannot delete Section {section.name} due to dependent records.")

    return redirect('manage_academic_structure')


@login_required
@user_passes_test(is_staff_member, login_url='login')
def edit_stream_view(request, pk):
    stream = get_object_or_404(Stream, pk=pk)
    if request.method == 'POST':
        name = request.POST.get('stream_name', '').strip()
        code = request.POST.get('stream_code', '').strip().upper()
        desc = request.POST.get('stream_desc', '').strip()

        conflict = Stream.objects.filter(code=code).exclude(pk=stream.pk).exists()
        if conflict:
            messages.error(request, f"Stream code '{code}' already in use.")
        else:
            stream.name = name
            stream.code = code
            stream.description = desc
            stream.save()
            messages.success(request, f"Stream '{stream.name}' updated.")

    return redirect('manage_academic_structure')


@login_required
@user_passes_test(is_staff_member, login_url='login')
def delete_stream_view(request, pk):
    stream = get_object_or_404(Stream, pk=pk)
    if stream.sections.exists():
        messages.error(request, f"Cannot delete stream '{stream.name}' while sections are assigned to it.")
    else:
        stream.delete()
        messages.success(request, f"Stream '{stream.name}' removed.")

    return redirect('manage_academic_structure')


# =====================================================================
# EXPORTS (EXCEL & PDF)
# =====================================================================

def get_all_filtered_enrollments(request):
    sessions = AcademicSession.objects.all().order_by('-start_date')
    selected_session_id = request.GET.get('session')

    if selected_session_id:
        selected_session = sessions.filter(id=selected_session_id).first()
    else:
        selected_session = sessions.filter(is_current=True).first() or sessions.first()

    selected_grade = request.GET.get('grade', '').strip()
    query = request.GET.get('q', '').strip()

    if not selected_session:
        return StudentEnrollment.objects.none()

    qs = (
        StudentEnrollment.objects
        .filter(session=selected_session)
        .select_related('student', 'section__academic_class', 'section__stream', 'session')
        .prefetch_related('enrolled_subjects__subject')
        .order_by('section__academic_class__numeric_level', 'section__name', 'roll_number')
    )

    if selected_grade:
        qs = qs.filter(section__academic_class__numeric_level=selected_grade)

    if query:
        qs = qs.filter(
            Q(student__student_name__icontains=query) |
            Q(student__admission_number__icontains=query) |
            Q(student__father_name__icontains=query) |
            Q(student__mother_name__icontains=query) |
            Q(student__mobile_number__icontains=query) |
            Q(roll_number__icontains=query)
        )
    return qs


@login_required
@user_passes_test(is_teacher_or_staff, login_url='login')
def export_students_excel(request):
    enrollments = get_all_filtered_enrollments(request)

    wb = Workbook()
    ws = wb.active
    ws.title = "Full Student Directory"

    header_fill = PatternFill(start_color="312E81", end_color="312E81", fill_type="solid")
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    align_center = Alignment(horizontal="center", vertical="center")

    headers = [
        "Session", "Class", "Section", "Stream", "Roll #",
        "Admission #", "Admission Date", "Student Name", "Father Name", "Mother Name",
        "Date of Birth", "Gender", "Category", "Mobile Number", "Aadhaar ID",
        "Permanent Address", "Bank Name", "Branch", "Account #", "IFSC Code",
        "Status", "Informed", "Informed At", "Enrolled Subjects"
    ]
    ws.append(headers)

    for col_idx in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=col_idx)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = align_center

    for enr in enrollments:
        student = enr.student
        stream_code = enr.section.stream.code if enr.section.stream else "-"
        subjects_str = ", ".join(es.subject.code or es.subject.name for es in enr.enrolled_subjects.all())

        ws.append([
            enr.session.name,
            enr.section.academic_class.name,
            enr.section.name,
            stream_code,
            enr.roll_number,
            student.admission_number,
            student.admission_date.strftime('%d/%m/%Y') if student.admission_date else "",
            student.student_name,
            student.father_name or "",
            student.mother_name or "",
            student.date_of_birth.strftime('%d/%m/%Y') if student.date_of_birth else "",
            student.get_gender_display(),
            student.get_caste_category_display(),
            student.mobile_number,
            "[Aadhaar Redacted]" if student.aadhaar_number else "",
            student.address or "",
            student.bank_name or "",
            student.bank_branch or "",
            student.account_number or "",
            student.ifsc_code or "",
            enr.get_status_display(),
            "Yes" if enr.is_informed else "Pending",
            enr.informed_at.strftime('%d/%m/%Y %H:%M') if enr.informed_at else "",
            subjects_str or "-"
        ])

    for col in ws.columns:
        max_len = max(len(str(cell.value or '')) for cell in col)
        col_letter = col[0].column_letter
        ws.column_dimensions[col_letter].width = max(max_len + 3, 12)

    response = HttpResponse(content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    response['Content-Disposition'] = 'attachment; filename="all_students_full_export.xlsx"'
    wb.save(response)
    return response


@login_required
@user_passes_test(is_teacher_or_staff, login_url='login')
def export_students_pdf(request):
    enrollments = get_all_filtered_enrollments(request)

    buffer = io.BytesIO()
    doc = SimpleDocTemplate(
        buffer,
        pagesize=landscape(letter),
        leftMargin=20,
        rightMargin=20,
        topMargin=25,
        bottomMargin=25
    )

    styles = getSampleStyleSheet()
    title_style = ParagraphStyle(
        'DocTitle',
        parent=styles['Heading1'],
        fontSize=15,
        leading=18,
        textColor=colors.HexColor('#1E293B'),
        spaceAfter=10
    )
    th_style = ParagraphStyle('THStyle', parent=styles['Normal'], fontSize=7, leading=9, fontName='Helvetica-Bold', textColor=colors.whitesmoke, alignment=1)
    td_style = ParagraphStyle('TDStyle', parent=styles['Normal'], fontSize=7, leading=8, textColor=colors.HexColor('#1E293B'))

    elements = [
        Paragraph("Complete Student Directory Export", title_style),
        Spacer(1, 4)
    ]

    table_data = [[
        Paragraph("Roll", th_style),
        Paragraph("Adm #", th_style),
        Paragraph("Name", th_style),
        Paragraph("Class/Sec", th_style),
        Paragraph("Parents (Father/Mother)", th_style),
        Paragraph("DOB / Gen / Cat", th_style),
        Paragraph("Contact & Address", th_style),
        Paragraph("Bank & A/C", th_style),
        Paragraph("Status", th_style),
        Paragraph("Informed", th_style),
        Paragraph("Subjects", th_style)
    ]]

    for enr in enrollments:
        s = enr.student
        sec_info = f"{enr.section.academic_class.name} ({enr.section.name})"
        if enr.section.stream:
            sec_info += f" - {enr.section.stream.code}"

        parents = f"F: {s.father_name or '-'}<br/>M: {s.mother_name or '-'}"
        demographics = f"{s.date_of_birth.strftime('%d/%m/%y') if s.date_of_birth else '-'}<br/>{s.get_gender_display()} | {s.get_caste_category_display()}"
        contact_info = f"{s.mobile_number}<br/>{s.address or '-'}"
        bank_info = f"{s.bank_name or '-'}<br/>{s.account_number or '-'}"
        subjects_str = ", ".join(es.subject.code or es.subject.name for es in enr.enrolled_subjects.all())

        table_data.append([
            Paragraph(str(enr.roll_number), td_style),
            Paragraph(s.admission_number, td_style),
            Paragraph(s.student_name, td_style),
            Paragraph(sec_info, td_style),
            Paragraph(parents, td_style),
            Paragraph(demographics, td_style),
            Paragraph(contact_info, td_style),
            Paragraph(bank_info, td_style),
            Paragraph(enr.get_status_display(), td_style),
            Paragraph("Informed" if enr.is_informed else "Pending", td_style),
            Paragraph(subjects_str or "-", td_style)
        ])

    table = Table(table_data, colWidths=[30, 45, 80, 75, 95, 75, 95, 80, 55, 45, 80])
    table.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#312E81')),
        ('GRID', (0, 0), (-1, -1), 0.5, colors.HexColor('#CBD5E1')),
        ('VALIGN', (0, 0), (-1, -1), 'TOP'),
        ('TOPPADDING', (0, 0), (-1, -1), 4),
        ('BOTTOMPADDING', (0, 0), (-1, -1), 4),
        ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, colors.HexColor('#F8FAFC')]),
    ]))

    elements.append(table)
    doc.build(elements)

    buffer.seek(0)
    response = HttpResponse(buffer, content_type='application/pdf')
    response['Content-Disposition'] = 'attachment; filename="all_students_full_export.pdf"'
    return response


@login_required
@user_passes_test(is_teacher_or_staff, login_url='login')
def export_single_student_excel(request, pk):
    student = get_object_or_404(Student, pk=pk)

    wb = Workbook()
    ws_profile = wb.active
    ws_profile.title = "Student Profile"

    primary_fill = PatternFill(start_color="4F46E5", end_color="4F46E5", fill_type="solid")
    white_bold = Font(name="Calibri", size=11, bold=True, color="FFFFFF")

    profile_data = [
        ("Field", "Details"),
        ("Admission Number", student.admission_number),
        ("Admission Date", student.admission_date.strftime('%d/%m/%Y') if student.admission_date else "-"),
        ("Student Name", student.student_name),
        ("Father's Name", student.father_name or "-"),
        ("Mother's Name", student.mother_name or "-"),
        ("Date of Birth", student.date_of_birth.strftime('%d/%m/%Y') if student.date_of_birth else "-"),
        ("Gender", student.get_gender_display()),
        ("Social Category", student.get_caste_category_display()),
        ("Mobile Number", student.mobile_number),
        ("Aadhaar Number", "[Aadhaar Redacted]" if student.aadhaar_number else "-"),
        ("Permanent Address", student.address or "-"),
        ("Bank Name", student.bank_name or "-"),
        ("Bank Branch", student.bank_branch or "-"),
        ("Account Number", student.account_number or "-"),
        ("IFSC Code", student.ifsc_code or "-"),
    ]

    for row in profile_data:
        ws_profile.append(row)

    ws_profile.cell(row=1, column=1).fill = primary_fill
    ws_profile.cell(row=1, column=1).font = white_bold
    ws_profile.cell(row=1, column=2).fill = primary_fill
    ws_profile.cell(row=1, column=2).font = white_bold

    ws_profile.column_dimensions['A'].width = 25
    ws_profile.column_dimensions['B'].width = 40

    response = HttpResponse(content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    response['Content-Disposition'] = f'attachment; filename="student_{student.admission_number}_profile.xlsx"'
    wb.save(response)
    return response


@login_required
@user_passes_test(is_teacher_or_staff, login_url='login')
def export_single_student_pdf(request, pk):
    student = get_object_or_404(Student, pk=pk)

    buffer = io.BytesIO()
    doc = SimpleDocTemplate(
        buffer,
        pagesize=portrait(letter),
        leftMargin=36,
        rightMargin=36,
        topMargin=36,
        bottomMargin=36
    )

    styles = getSampleStyleSheet()
    title_style = ParagraphStyle(
        'DocTitle',
        parent=styles['Heading1'],
        fontSize=18,
        leading=22,
        textColor=colors.HexColor('#1E293B'),
        spaceAfter=4
    )
    sub_title = ParagraphStyle('SubTitle', parent=styles['Normal'], fontSize=10, textColor=colors.HexColor('#64748B'), spaceAfter=15)
    section_heading = ParagraphStyle('SectionHeading', parent=styles['Heading3'], fontSize=12, leading=15, textColor=colors.HexColor('#4338CA'), spaceBefore=12, spaceAfter=6)
    cell_bold = ParagraphStyle('CellBold', parent=styles['Normal'], fontSize=9, fontName='Helvetica-Bold')
    cell_regular = ParagraphStyle('CellReg', parent=styles['Normal'], fontSize=9)

    elements = [
        Paragraph("Student Profile Dossier", title_style),
        Paragraph(f"Generated for: <b>{student.student_name}</b> (Adm: {student.admission_number})", sub_title),
    ]

    # Section 1: Personal Details
    elements.append(Paragraph("1. Personal & Contact Information", section_heading))
    personal_data = [
        [Paragraph("Full Name", cell_bold), Paragraph(student.student_name, cell_regular), Paragraph("Admission Date", cell_bold), Paragraph(student.admission_date.strftime('%d/%m/%Y') if student.admission_date else "-", cell_regular)],
        [Paragraph("Father's Name", cell_bold), Paragraph(student.father_name or "-", cell_regular), Paragraph("Mother's Name", cell_bold), Paragraph(student.mother_name or "-", cell_regular)],
        [Paragraph("Date of Birth", cell_bold), Paragraph(student.date_of_birth.strftime('%d/%m/%Y') if student.date_of_birth else "-", cell_regular), Paragraph("Gender", cell_bold), Paragraph(student.get_gender_display(), cell_regular)],
        [Paragraph("Category", cell_bold), Paragraph(student.get_caste_category_display(), cell_regular), Paragraph("Contact No.", cell_bold), Paragraph(student.mobile_number, cell_regular)],
        [Paragraph("Aadhaar ID", cell_bold), Paragraph("[Aadhaar Redacted]" if student.aadhaar_number else "-", cell_regular), Paragraph("Address", cell_bold), Paragraph(student.address or "-", cell_regular)],
    ]
    t_personal = Table(personal_data, colWidths=[100, 170, 100, 170])
    t_personal.setStyle(TableStyle([
        ('GRID', (0, 0), (-1, -1), 0.5, colors.HexColor('#CBD5E1')),
        ('BACKGROUND', (0, 0), (0, -1), colors.HexColor('#F8FAFC')),
        ('BACKGROUND', (2, 0), (2, -1), colors.HexColor('#F8FAFC')),
        ('VALIGN', (0, 0), (-1, -1), 'TOP'),
        ('TOPPADDING', (0, 0), (-1, -1), 5),
        ('BOTTOMPADDING', (0, 0), (-1, -1), 5),
    ]))
    elements.append(t_personal)

    # Section 2: Banking
    elements.append(Paragraph("2. Direct Benefit & Bank Credentials", section_heading))
    bank_data = [
        [Paragraph("Bank Name", cell_bold), Paragraph(student.bank_name or "-", cell_regular), Paragraph("Branch", cell_bold), Paragraph(student.bank_branch or "-", cell_regular)],
        [Paragraph("Account No.", cell_bold), Paragraph(student.account_number or "-", cell_regular), Paragraph("IFSC Code", cell_bold), Paragraph(student.ifsc_code or "-", cell_regular)],
    ]
    t_bank = Table(bank_data, colWidths=[100, 170, 100, 170])
    t_bank.setStyle(TableStyle([
        ('GRID', (0, 0), (-1, -1), 0.5, colors.HexColor('#CBD5E1')),
        ('BACKGROUND', (0, 0), (0, -1), colors.HexColor('#F8FAFC')),
        ('BACKGROUND', (2, 0), (2, -1), colors.HexColor('#F8FAFC')),
        ('VALIGN', (0, 0), (-1, -1), 'TOP'),
        ('TOPPADDING', (0, 0), (-1, -1), 5),
        ('BOTTOMPADDING', (0, 0), (-1, -1), 5),
    ]))
    elements.append(t_bank)

    # Section 3: Academic History
    elements.append(Paragraph("3. Academic Enrollment & Subjects", section_heading))
    academic_table_data = [[
        Paragraph("Session", cell_bold),
        Paragraph("Class & Section", cell_bold),
        Paragraph("Roll #", cell_bold),
        Paragraph("Status", cell_bold),
        Paragraph("Enrolled Subjects", cell_bold)
    ]]

    enrollments = student.enrollments.select_related(
        'session', 'section__academic_class', 'section__stream'
    ).prefetch_related('enrolled_subjects__subject')

    for enr in enrollments:
        sec_name = f"{enr.section.academic_class.name} ({enr.section.name})"
        if enr.section.stream:
            sec_name += f" - {enr.section.stream.code}"
        subj_list = ", ".join(es.subject.code or es.subject.name for es in enr.enrolled_subjects.all())

        academic_table_data.append([
            Paragraph(enr.session.name, cell_regular),
            Paragraph(sec_name, cell_regular),
            Paragraph(str(enr.roll_number), cell_regular),
            Paragraph(enr.get_status_display(), cell_regular),
            Paragraph(subj_list or "-", cell_regular),
        ])

    t_academic = Table(academic_table_data, colWidths=[90, 140, 50, 90, 170])
    t_academic.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#4F46E5')),
        ('TEXTCOLOR', (0, 0), (-1, 0), colors.whitesmoke),
        ('GRID', (0, 0), (-1, -1), 0.5, colors.HexColor('#CBD5E1')),
        ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
        ('TOPPADDING', (0, 0), (-1, -1), 5),
        ('BOTTOMPADDING', (0, 0), (-1, -1), 5),
    ]))
    elements.append(t_academic)

    doc.build(elements)
    buffer.seek(0)
    response = HttpResponse(buffer, content_type='application/pdf')
    response['Content-Disposition'] = f'attachment; filename="student_{student.admission_number}_dossier.pdf"'
    return response